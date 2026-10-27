"""Vrais fichiers : Word (.docx), PDF, PowerPoint (.pptx), Excel (.xlsx) — et modification d'un Word ou d'un Excel déposé.

Le modèle écrit le contenu en Markdown simple (titres #, listes -, 1., tableaux |, **gras**, *italique*, --- = saut de page) ;
pour Excel, des feuilles de lignes (ou les tableaux du Markdown). Chaque fichier produit est rouvert pour vérification
(pages, diapositives, lignes) avant d'être proposé au téléchargement. Rien n'écrase un fichier existant."""
import datetime as dt
import re
from pathlib import Path

FORMATS = {"docx", "pdf", "pptx", "xlsx"}
_MAX_TEXTE = 200_000


class DocumentError(RuntimeError):
    pass


# ---------------------------------------------------------------- Markdown → blocs

def blocs(md: str) -> list[dict]:
    """Découpe un Markdown simple en blocs : titre, para, liste, tableau, citation, saut, code."""
    lignes = (md or "").replace("\r\n", "\n").split("\n")
    out: list[dict] = []
    para: list[str] = []
    i = 0

    def vider():
        if para:
            out.append({"t": "para", "texte": " ".join(x.strip() for x in para)})
            para.clear()

    while i < len(lignes):
        l = lignes[i]
        s = l.strip()
        if s.startswith("```"):
            vider()
            code = []
            i += 1
            while i < len(lignes) and not lignes[i].strip().startswith("```"):
                code.append(lignes[i])
                i += 1
            out.append({"t": "code", "texte": "\n".join(code)})
        elif not s:
            vider()
        elif m := re.match(r"^(#{1,4})\s+(.*)$", s):
            vider()
            out.append({"t": "titre", "niveau": len(m[1]), "texte": m[2].strip()})
        elif re.fullmatch(r"-{3,}|\*{3,}|_{3,}", s):
            vider()
            out.append({"t": "saut"})
        elif s.startswith("|"):
            vider()
            rangs = []
            while i < len(lignes) and lignes[i].strip().startswith("|"):
                cellules = [c.strip() for c in lignes[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cellules if c):
                    rangs.append(cellules)
                i += 1
            if rangs:
                n = max(len(r) for r in rangs)
                out.append({"t": "tableau", "lignes": [r + [""] * (n - len(r)) for r in rangs]})
            continue
        elif re.match(r"^([-*+•]|\d+[.)])\s+", s):
            vider()
            numerote = bool(re.match(r"^\d+[.)]", s))
            items = []
            while i < len(lignes) and re.match(r"^\s*([-*+•]|\d+[.)])\s+", lignes[i]):
                brut = lignes[i]
                niveau = min(2, (len(brut) - len(brut.lstrip(" "))) // 2)
                items.append({"texte": re.sub(r"^\s*([-*+•]|\d+[.)])\s+", "", brut).strip(), "niveau": niveau})
                i += 1
            out.append({"t": "liste", "numerote": numerote, "items": items})
            continue
        elif s.startswith(">"):
            vider()
            out.append({"t": "citation", "texte": s.lstrip("> ").strip()})
        else:
            para.append(l)
        i += 1
    vider()
    return out


_INLINE = re.compile(r"(\*\*[^*]+\*\*|__[^_]+__|\*[^*\s][^*]*\*|_[^_\s][^_]*_|`[^`]+`)")


def morceaux(texte: str) -> list[tuple[str, str]]:
    """« du **gras** et *ça* » → [("", "du "), ("b", "gras"), ("", " et "), ("i", "ça")] ; liens [t](u) → « t (u) »."""
    texte = re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)", r"\1 (\2)", texte or "")
    out = []
    for p in _INLINE.split(texte):
        if not p:
            continue
        if p.startswith(("**", "__")) and len(p) > 4:
            out.append(("b", p[2:-2]))
        elif p.startswith("`"):
            out.append(("c", p[1:-1]))
        elif p[0] in "*_" and p[-1] == p[0] and len(p) > 2:
            out.append(("i", p[1:-1]))
        else:
            out.append(("", p))
    return out


def brut(texte: str) -> str:
    return "".join(t for _k, t in morceaux(texte))


# ---------------------------------------------------------------- Word

def _docx(titre: str, md: str, chemin: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)
    if titre:
        doc.add_heading(titre, level=0)
        doc.core_properties.title = titre
    doc.core_properties.author = "EVA"

    def remplir(par, texte):
        for k, t in morceaux(texte):
            r = par.add_run(t)
            r.bold, r.italic = k == "b", k == "i"
            if k == "c":
                r.font.name = "Consolas"
    for b in blocs(md):
        if b["t"] == "titre":
            if titre and b["niveau"] == 1 and b["texte"] == titre:
                continue
            doc.add_heading(brut(b["texte"]), level=min(b["niveau"], 4))
        elif b["t"] == "para":
            remplir(doc.add_paragraph(), b["texte"])
        elif b["t"] == "citation":
            remplir(doc.add_paragraph(style="Quote"), b["texte"])
        elif b["t"] == "code":
            p = doc.add_paragraph()
            r = p.add_run(b["texte"])
            r.font.name, r.font.size = "Consolas", Pt(9)
        elif b["t"] == "liste":
            for it in b["items"]:
                style = ("List Number" if b["numerote"] else "List Bullet") + (f" {it['niveau'] + 1}" if it["niveau"] else "")
                remplir(doc.add_paragraph(style=style), it["texte"])
        elif b["t"] == "tableau":
            rangs = b["lignes"]
            t = doc.add_table(rows=len(rangs), cols=len(rangs[0]))
            t.style = "Light Grid Accent 1"
            for i, r in enumerate(rangs):
                for j, c in enumerate(r):
                    cell = t.cell(i, j)
                    cell.text = ""
                    remplir(cell.paragraphs[0], c)
                    if i == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            doc.add_paragraph()
        elif b["t"] == "saut":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.save(chemin)


# ---------------------------------------------------------------- PDF

def _police_pdf() -> tuple[str, str]:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    import reportlab
    rl = Path(reportlab.__file__).parent / "fonts"
    for d, n, g in (("/usr/share/fonts/truetype/dejavu", "DejaVuSans.ttf", "DejaVuSans-Bold.ttf"),
                    ("/usr/share/fonts/dejavu", "DejaVuSans.ttf", "DejaVuSans-Bold.ttf"),
                    (str(rl), "Vera.ttf", "VeraBd.ttf")):                 # Vera : livrée avec reportlab (accents, €)
        n, g = Path(d) / n, Path(d) / g
        if n.is_file() and g.is_file():
            try:
                pdfmetrics.registerFont(TTFont("EvaSans", str(n)))
                pdfmetrics.registerFont(TTFont("EvaSans-Bold", str(g)))
                return "EvaSans", "EvaSans-Bold"
            except Exception:  # noqa: BLE001
                break
    return "Helvetica", "Helvetica-Bold"


def _pdf(titre: str, md: str, chemin: Path) -> None:
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import ListFlowable, ListItem, PageBreak, Paragraph, Preformatted, SimpleDocTemplate, Spacer, Table, TableStyle

    police, gras = _police_pdf()
    import reportlab
    vera = police == "EvaSans" and not Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf").is_file() and not Path("/usr/share/fonts/dejavu/DejaVuSans.ttf").is_file()
    _ = reportlab

    def sur(t: str) -> str:
        """Police réduite (Vera ou Helvetica) : flèches et coches remplacées, émojis retirés (sinon des carrés vides)."""
        if police == "EvaSans" and not vera:
            return t
        for a, b in (("→", "->"), ("←", "<-"), ("⇒", "=>"), ("✓", "v"), ("✔", "v"), ("✗", "x"), ("✕", "x"), ("•", "-"), ("…", "..."),
                     ("≤", "<="), ("≥", ">="), ("≠", "!="), ("−", "-"), ("\u202f", " ")):
            t = t.replace(a, b)
        return "".join(c for c in t if ord(c) < 0x2000 or c in "€–—‘’“”‰")
    ss = getSampleStyleSheet()
    corps = ParagraphStyle("corps", parent=ss["BodyText"], fontName=police, fontSize=10.5, leading=14.5, spaceAfter=6)
    titres = {n: ParagraphStyle(f"t{n}", parent=ss[f"Heading{n}"], fontName=gras, textColor=colors.HexColor("#1f3a68")) for n in (1, 2, 3, 4)}
    grand = ParagraphStyle("titre", parent=ss["Title"], fontName=gras, textColor=colors.HexColor("#1f3a68"))
    cit = ParagraphStyle("cit", parent=corps, leftIndent=14, textColor=colors.HexColor("#555555"), fontName=police)
    code = ParagraphStyle("code", parent=corps, fontName="Courier", fontSize=8.5, leading=11, backColor=colors.HexColor("#f3f3f3"))

    def html(texte):
        out = []
        for k, t in morceaux(sur(texte)):
            t = escape(t)
            out.append(f"<b>{t}</b>" if k == "b" else f"<i>{t}</i>" if k == "i" else f'<font face="Courier">{t}</font>' if k == "c" else t)
        return "".join(out)

    flux = []
    if titre:
        flux += [Paragraph(escape(sur(titre)), grand), Spacer(1, 6)]
    for b in blocs(md):
        if b["t"] == "titre":
            if titre and b["niveau"] == 1 and b["texte"] == titre:
                continue
            flux.append(Paragraph(html(b["texte"]), titres[min(b["niveau"], 4)]))
        elif b["t"] == "para":
            flux.append(Paragraph(html(b["texte"]), corps))
        elif b["t"] == "citation":
            flux.append(Paragraph(html(b["texte"]), cit))
        elif b["t"] == "code":
            flux.append(Preformatted(sur(b["texte"]), code))
        elif b["t"] == "liste":
            flux.append(ListFlowable([ListItem(Paragraph(html(it["texte"]), corps), leftIndent=12 + 12 * it["niveau"]) for it in b["items"]],
                                     bulletType="1" if b["numerote"] else "bullet", start="1" if b["numerote"] else None,
                                     bulletFontName=police, leftIndent=14))
        elif b["t"] == "tableau":
            cellule = ParagraphStyle("cell", parent=corps, fontSize=9, leading=11.5, spaceAfter=0)
            entete = ParagraphStyle("ent", parent=cellule, fontName=gras, textColor=colors.white)
            data = [[Paragraph(html(c), entete if i == 0 else cellule) for c in r] for i, r in enumerate(b["lignes"])]
            t = Table(data, repeatRows=1, colWidths=[(A4[0] - 4 * cm) / len(b["lignes"][0])] * len(b["lignes"][0]))
            t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f3a68")),
                                   ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#eef2f8")]),
                                   ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#b8c4d6")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            flux += [t, Spacer(1, 8)]
        elif b["t"] == "saut":
            flux.append(PageBreak())

    def pied(canvas, d):
        canvas.saveState()
        canvas.setFont(police, 8)
        canvas.setFillColor(colors.HexColor("#888888"))
        canvas.drawRightString(A4[0] - 2 * cm, 1.2 * cm, f"{sur(titre)[:80] + ' — ' if titre else ''}page {d.page}")
        canvas.restoreState()
    SimpleDocTemplate(str(chemin), pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm, bottomMargin=2 * cm,
                      title=titre or chemin.stem, author="EVA").build(flux, onFirstPage=pied, onLaterPages=pied)


# ---------------------------------------------------------------- PowerPoint

def _pptx(titre: str, md: str, chemin: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    bl = blocs(md)
    sous_titre = ""
    if bl and bl[0]["t"] == "titre" and bl[0]["niveau"] == 1:
        titre = titre or brut(bl[0]["texte"])
        bl = bl[1:]
        if bl and bl[0]["t"] == "para":
            sous_titre, bl = brut(bl[0]["texte"]), bl[1:]
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = titre or chemin.stem
    s.placeholders[1].text = sous_titre or dt.date.today().strftime("%d/%m/%Y")

    diapos: list[dict] = []
    for b in bl:
        if b["t"] == "titre" and b["niveau"] <= 2:
            diapos.append({"titre": brut(b["texte"]), "contenu": []})
        elif b["t"] != "saut":
            if not diapos:
                diapos.append({"titre": titre, "contenu": []})
            diapos[-1]["contenu"].append(b)

    def puces(d):
        out = []
        for b in d["contenu"]:
            if b["t"] == "liste":
                out += [(it["texte"], it["niveau"]) for it in b["items"]]
            elif b["t"] in ("para", "citation"):
                out.append((b["texte"], 0))
            elif b["t"] == "titre":
                out.append((f"**{brut(b['texte'])}**", 0))
            elif b["t"] == "code":
                out += [(f"`{x}`", 1) for x in b["texte"].splitlines()[:8]]
        return out

    for d in diapos:
        tableau = next((b for b in d["contenu"] if b["t"] == "tableau"), None)
        lignes = puces(d)
        paquets = [lignes[i:i + 7] for i in range(0, len(lignes), 7)] or [[]]
        for n, paquet in enumerate(paquets):
            s = prs.slides.add_slide(prs.slide_layouts[5 if tableau and n == 0 and not paquet else 1])
            s.shapes.title.text = d["titre"] + (" (suite)" if n else "")
            if paquet and len(s.placeholders) > 1:
                tf = s.placeholders[1].text_frame
                tf.clear()
                for k, (texte, niveau) in enumerate(paquet):
                    p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                    p.level = niveau
                    for genre, t in morceaux(texte):
                        r = p.add_run()
                        r.text = t
                        r.font.bold, r.font.italic = genre == "b", genre == "i"
                        r.font.size = Pt(24 if niveau == 0 else 20)
            if tableau and n == 0:
                rangs = tableau["lignes"][:15]
                haut = Inches(1.6 if not paquet else 4.6)
                if paquet:
                    s = prs.slides.add_slide(prs.slide_layouts[5])
                    s.shapes.title.text = d["titre"]
                    haut = Inches(1.6)
                forme = s.shapes.add_table(len(rangs), len(rangs[0]), Inches(0.6), haut, Inches(12.1), Inches(0.4) * len(rangs))
                for i, r in enumerate(rangs):
                    for j, c in enumerate(r):
                        cell = forme.table.cell(i, j)
                        cell.text = brut(c)
                        for p in cell.text_frame.paragraphs:
                            for r_ in p.runs:
                                r_.font.size = Pt(14)
                                r_.font.bold = i == 0
    prs.core_properties.title = titre or chemin.stem
    prs.core_properties.author = "EVA"
    prs.save(chemin)


# ---------------------------------------------------------------- Excel

_NOMBRE = re.compile(r"^-?[\d\s  ]*\d([.,]\d+)?$")


def valeur(c):
    """Texte d'une cellule → nombre, pourcentage, montant, date ou formule quand c'est évident."""
    if not isinstance(c, str):
        return c, None
    s = brut(c).strip()
    if not s:
        return None, None
    if s.startswith("="):
        return s, None
    if m := re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})", s):
        try:
            return dt.date(int(m[1]), int(m[2]), int(m[3])), "DD/MM/YYYY"
        except ValueError:
            return s, None
    if m := re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", s):
        try:
            return dt.date(int(m[3]), int(m[2]), int(m[1])), "DD/MM/YYYY"
        except ValueError:
            return s, None
    fmt = None
    t = s
    if t.endswith("%"):
        t, fmt = t[:-1].strip(), "0.0%"
    elif t.endswith("€") or t.startswith("€"):
        t, fmt = t.strip("€ ").strip(), '#,##0.00 "€"'
    if _NOMBRE.match(t) and re.search(r"\d", t):
        n = t.replace(" ", "").replace(" ", "").replace(" ", "")
        if n.count(",") == 1 and "." not in n:
            n = n.replace(",", ".")
        try:
            v = float(n)
        except ValueError:
            return s, None
        if fmt == "0.0%":
            v /= 100
        elif v.is_integer() and fmt is None and not re.match(r"^0\d", t):
            v = int(v)
        elif fmt is None and not v.is_integer():
            fmt = "#,##0.00"
        return v, fmt
    return s, None


def _xlsx(titre: str, feuilles: list[dict], chemin: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    noms = set()
    for k, f in enumerate(feuilles, 1):
        nom = re.sub(r"[\\/*?:\[\]]", " ", str(f.get("nom") or f"Feuille {k}"))[:31].strip() or f"Feuille {k}"
        while nom in noms:
            nom = (nom[:28] + f" {k}")[:31]
        noms.add(nom)
        ws = wb.create_sheet(nom)
        lignes = [list(r) if isinstance(r, (list, tuple)) else [r] for r in (f.get("lignes") or [])][:50000]
        for i, r in enumerate(lignes, 1):
            for j, c in enumerate(r[:200], 1):
                v, fmt = valeur(c) if i > 1 or not f.get("entete", True) else (brut(str(c)) if isinstance(c, str) else c, None)
                cell = ws.cell(row=i, column=j, value=v)
                if fmt:
                    cell.number_format = fmt
        if lignes and f.get("entete", True):
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F3A68")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            ws.freeze_panes = "A2"
            if len(lignes) > 1:
                ws.auto_filter.ref = ws.dimensions
        for j in range(1, (max((len(r) for r in lignes), default=0)) + 1):
            larg = max((len(str(ws.cell(row=i, column=j).value or "")) for i in range(1, min(len(lignes), 500) + 1)), default=8)
            ws.column_dimensions[get_column_letter(j)].width = max(8, min(60, larg + 2))
    wb.properties.title = titre or chemin.stem
    wb.properties.creator = "EVA"
    wb.save(chemin)


def feuilles_depuis_md(md: str) -> list[dict]:
    out = []
    titre_courant = ""
    for b in blocs(md):
        if b["t"] == "titre":
            titre_courant = brut(b["texte"])
        elif b["t"] == "tableau":
            out.append({"nom": titre_courant or f"Tableau {len(out) + 1}", "lignes": b["lignes"]})
    return out


# ---------------------------------------------------------------- vérification

def verifier(chemin: Path) -> dict:
    """Rouvre le fichier produit : il doit se lire, avec le contenu attendu."""
    ext = chemin.suffix.lower()
    try:
        if ext == ".docx":
            from docx import Document
            d = Document(chemin)
            texte = " ".join(p.text for p in d.paragraphs)
            return {"ok": bool(texte.strip() or d.tables), "paragraphes": len(d.paragraphs), "tableaux": len(d.tables), "mots": len(texte.split())}
        if ext == ".pdf":
            from pypdf import PdfReader
            r = PdfReader(str(chemin))
            texte = "".join((p.extract_text() or "") for p in r.pages[:3])
            return {"ok": len(r.pages) > 0 and bool(texte.strip()), "pages": len(r.pages)}
        if ext == ".pptx":
            from pptx import Presentation
            p = Presentation(chemin)
            return {"ok": len(p.slides) > 0, "diapositives": len(p.slides)}
        if ext == ".xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(chemin)
            return {"ok": True, "feuilles": [{"nom": ws.title, "lignes": ws.max_row, "colonnes": ws.max_column} for ws in wb.worksheets]}
    except Exception as exc:  # noqa: BLE001
        return {"ok": False, "erreur": f"Fichier illisible après création : {exc}"[:200]}
    return {"ok": chemin.is_file()}


# ---------------------------------------------------------------- points d'entrée

def nom_libre(dossier: Path, nom: str, ext: str) -> Path:
    base = re.sub(r"[^\w\- .()àâäéèêëîïôöùûüçœÀÂÉÈÊÎÔÙÛÇ]", "_", Path(nom or "document").stem).strip(" ._")[:80] or "document"
    p = dossier / f"{base}.{ext}"
    n = 2
    while p.exists():
        p = dossier / f"{base} ({n}).{ext}"
        n += 1
    return p


def creer(dossier: Path, format_: str, nom: str, titre: str = "", contenu: str = "", feuilles: list | None = None) -> Path:
    format_ = (format_ or "").lower().strip(". ")
    format_ = {"word": "docx", "doc": "docx", "excel": "xlsx", "xls": "xlsx", "powerpoint": "pptx", "ppt": "pptx", "diaporama": "pptx"}.get(format_, format_)
    if format_ not in FORMATS:
        raise DocumentError("Format inconnu : docx (Word), pdf, pptx (PowerPoint) ou xlsx (Excel).")
    if len(contenu or "") > _MAX_TEXTE:
        raise DocumentError("Contenu trop long (200 000 caractères au plus).")
    dossier.mkdir(parents=True, exist_ok=True)
    chemin = nom_libre(dossier, nom or titre, format_)
    if format_ == "xlsx":
        fs = [f for f in (feuilles or []) if isinstance(f, dict)] or feuilles_depuis_md(contenu)
        if not fs:
            raise DocumentError("Pour Excel, donne des feuilles (nom + lignes) ou des tableaux Markdown.")
        _xlsx(titre, fs, chemin)
    else:
        if not (contenu or "").strip():
            raise DocumentError("Contenu vide.")
        {"docx": _docx, "pdf": _pdf, "pptx": _pptx}[format_](titre, contenu, chemin)
    return chemin


def modifier(source: Path, dossier: Path, remplacements: list | None = None, ajout: str = "", cellules: list | None = None,
             lignes: list | None = None, feuille: str = "") -> Path:
    """Copie modifiée d'un Word ou d'un Excel (l'original n'est jamais touché)."""
    ext = source.suffix.lower()
    dossier.mkdir(parents=True, exist_ok=True)
    cible = nom_libre(dossier, source.stem + " (modifié)", ext.lstrip("."))
    remplacements = [r for r in (remplacements or []) if isinstance(r, dict) and r.get("chercher")]
    if ext == ".docx":
        from docx import Document
        d = Document(source)
        trouves = 0
        paragraphes = list(d.paragraphs) + [p for t in d.tables for row in t.rows for c in row.cells for p in c.paragraphs]
        for r in remplacements:
            avant, apres = str(r["chercher"]), str(r.get("remplacer", ""))
            for p in paragraphes:
                if avant in p.text:
                    trouves += p.text.count(avant)
                    if any(avant in run.text for run in p.runs):
                        for run in p.runs:
                            run.text = run.text.replace(avant, apres)
                    else:                                    # texte coupé entre plusieurs morceaux : on garde le style du premier
                        t = p.text.replace(avant, apres)
                        for k, run in enumerate(p.runs):
                            run.text = t if k == 0 else ""
        if remplacements and not trouves:
            raise DocumentError("Aucun des textes à remplacer n'a été trouvé dans le document.")
        if ajout.strip():
            for b in blocs(ajout):
                if b["t"] == "titre":
                    d.add_heading(brut(b["texte"]), level=min(b["niveau"], 4))
                elif b["t"] == "liste":
                    for it in b["items"]:
                        d.add_paragraph(brut(it["texte"]), style="List Number" if b["numerote"] else "List Bullet")
                elif b["t"] == "tableau":
                    t = d.add_table(rows=len(b["lignes"]), cols=len(b["lignes"][0]))
                    t.style = "Table Grid"
                    for i, row in enumerate(b["lignes"]):
                        for j, c in enumerate(row):
                            t.cell(i, j).text = brut(c)
                elif b["t"] in ("para", "citation", "code"):
                    d.add_paragraph(brut(b["texte"]))
        d.save(cible)
        return cible
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(source)
        ws = wb[feuille] if feuille and feuille in wb.sheetnames else wb.active
        for c in cellules or []:
            if not isinstance(c, dict) or not re.fullmatch(r"[A-Za-z]{1,3}\d{1,7}", str(c.get("cellule", ""))):
                raise DocumentError(f"Cellule invalide : {c}")
            cible_ws = wb[c["feuille"]] if c.get("feuille") in wb.sheetnames else ws
            v, fmt = valeur(c.get("valeur")) if isinstance(c.get("valeur"), str) else (c.get("valeur"), None)
            cible_ws[str(c["cellule"]).upper()] = v
            if fmt:
                cible_ws[str(c["cellule"]).upper()].number_format = fmt
        for r in lignes or []:
            ws.append([valeur(x)[0] for x in (r if isinstance(r, list) else [r])])
        for r in remplacements:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and str(r["chercher"]) in cell.value:
                        cell.value = cell.value.replace(str(r["chercher"]), str(r.get("remplacer", "")))
        wb.save(cible.with_suffix(ext))
        return cible.with_suffix(ext)
    raise DocumentError("Je sais modifier un Word (.docx) ou un Excel (.xlsx). Pour un PDF ou un PowerPoint, je crée une nouvelle version.")
